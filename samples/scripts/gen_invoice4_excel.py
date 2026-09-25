import openpyxl
from openpyxl.styles import Font, Alignment

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Invoice"

# Title / metadata block scattered above the table, not a clean header-row-1 layout
ws["A1"] = "MERIDIAN TRADING CO."
ws["A1"].font = Font(bold=True, size=16)
ws["A2"] = "Wholesale Import / Export"
ws["A4"] = "Statement of Charges"
ws["A4"].font = Font(bold=True, size=12)

ws["A6"] = "Ref No:"
ws["B6"] = "MTC-30845"
ws["A7"] = "Issued:"
ws["B7"] = "2026-05-09"
ws["A8"] = "Client:"
ws["B8"] = "Harborview Distribution LLC"

# blank row 9, table starts at row 11 with unusual header names
ws["A11"] = "Item"
ws["B11"] = "Units"
ws["C11"] = "Price/Unit"
ws["D11"] = "Line Amount"
for col in ("A11", "B11", "C11", "D11"):
    ws[col].font = Font(bold=True)

rows = [
    ("Imported Ceramic Tile - 12x12 (box of 10)", 60, 22.50, 1350.00),
    ("Packing Crate, Reinforced", 25, 18.00, 450.00),
    ("Freight Handling Surcharge", 1, 200.00, 200.00),
]
r = 12
for desc, qty, price, amt in rows:
    ws.cell(row=r, column=1, value=desc)
    ws.cell(row=r, column=2, value=qty)
    ws.cell(row=r, column=3, value=price)
    ws.cell(row=r, column=4, value=amt)
    r += 1

# totals placed a couple rows below with a gap, in a different column arrangement
r += 1
ws.cell(row=r, column=3, value="Subtotal").font = Font(bold=True)
ws.cell(row=r, column=4, value=2000.00)
r += 1
ws.cell(row=r, column=3, value="Handling Fee").font = Font(bold=True)
ws.cell(row=r, column=4, value=0.00)
r += 1
ws.cell(row=r, column=3, value="AMOUNT DUE").font = Font(bold=True)
ws.cell(row=r, column=4, value=2000.00).font = Font(bold=True)

ws.column_dimensions["A"].width = 42
ws.column_dimensions["B"].width = 10
ws.column_dimensions["C"].width = 14
ws.column_dimensions["D"].width = 14

wb.save("samples/input/invoice-04-excel.xlsx")
print("wrote samples/input/invoice-04-excel.xlsx")
